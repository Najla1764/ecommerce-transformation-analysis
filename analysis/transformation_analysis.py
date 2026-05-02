# ============================================
# E-Commerce Business Transformation Analysis
# Author: Najla Z
# Tool: Python + Excel + Dashboard
# ============================================

import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
import os

# ============================================
# STEP 1 — LOAD THE DATA
# ============================================

print("Loading data...")
df = pd.read_excel("data/ecommerce_data.xlsx", header=1)
df.columns = df.columns.str.strip()
print("Columns found:", df.columns.tolist())

print(f"Total records loaded: {len(df)}")
print(f"Columns: {list(df.columns)}")
print("")

# ============================================
# STEP 2 — ANALYZE THE DATA
# ============================================

print("Analyzing processes...")

# Group by process and calculate averages
analysis = df.groupby("process").agg(
    total_manual_hours = ("manual_time_hrs", "sum"),
    avg_manual_hours   = ("manual_time_hrs", "mean"),
    avg_error_rate     = ("error_rate_%", "mean"),
    record_count       = ("order_id", "count")
).reset_index()

# Calculate ROI Score
# Formula: higher hours + higher error rate = higher ROI if automated
analysis["roi_score"] = (
    (analysis["avg_manual_hours"] * 0.5) +
    (analysis["avg_error_rate"] * 0.5)
).round(2)

# Rank by ROI Score
analysis["rank"] = analysis["roi_score"].rank(ascending=False).astype(int)

# Add automation label
def get_automation_level(score):
    if score >= 10:
        return "HIGH"
    elif score >= 7:
        return "MEDIUM"
    else:
        return "LOW"

analysis["automation_priority"] = analysis["roi_score"].apply(get_automation_level)

# Sort by rank
analysis = analysis.sort_values("rank")

print("Analysis complete!")
print("")
print(analysis[["process", "roi_score", "rank", "automation_priority"]])
print("")

# ============================================
# STEP 3 — GENERATE EXCEL REPORT
# ============================================

print("Generating Excel report...")

wb = Workbook()

# --- SHEET 1: SUMMARY ---
ws1 = wb.active
ws1.title = "Summary"

# Title
ws1["A1"] = "E-Commerce Business Transformation Analysis — Summary"
ws1["A1"].font = Font(bold=True, size=14, color="FFFFFF")
ws1["A1"].fill = PatternFill("solid", fgColor="0F3460")
ws1["A1"].alignment = Alignment(horizontal="center")
ws1.merge_cells("A1:F1")

# Subtitle
ws1["A2"] = "Top 5 Processes Recommended for Immediate Automation"
ws1["A2"].font = Font(bold=True, size=11, color="0F3460")
ws1.merge_cells("A2:F2")

# Headers
headers = ["Rank", "Process", "Avg Manual Hours", "Avg Error Rate %", "ROI Score", "Automation Priority"]
for col, header in enumerate(headers, 1):
    cell = ws1.cell(row=3, column=col, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="16213E")
    cell.alignment = Alignment(horizontal="center")

# Top 5 rows
top5 = analysis.head(5)
for row_idx, (_, row) in enumerate(top5.iterrows(), 4):
    ws1.cell(row=row_idx, column=1, value=int(row["rank"]))
    ws1.cell(row=row_idx, column=2, value=row["process"])
    ws1.cell(row=row_idx, column=3, value=round(row["avg_manual_hours"], 2))
    ws1.cell(row=row_idx, column=4, value=round(row["avg_error_rate"], 2))
    ws1.cell(row=row_idx, column=5, value=row["roi_score"])
    ws1.cell(row=row_idx, column=6, value=row["automation_priority"])
    # Color HIGH rows
    if row["automation_priority"] == "HIGH":
        for col in range(1, 7):
            ws1.cell(row=row_idx, column=col).fill = PatternFill("solid", fgColor="E8EAF6")

# Column widths
ws1.column_dimensions["A"].width = 8
ws1.column_dimensions["B"].width = 28
ws1.column_dimensions["C"].width = 20
ws1.column_dimensions["D"].width = 18
ws1.column_dimensions["E"].width = 12
ws1.column_dimensions["F"].width = 22

# --- SHEET 2: FULL ANALYSIS ---
ws2 = wb.create_sheet("Full Analysis")

ws2["A1"] = "Full Process Analysis — All 10 Processes"
ws2["A1"].font = Font(bold=True, size=13, color="FFFFFF")
ws2["A1"].fill = PatternFill("solid", fgColor="0F3460")
ws2.merge_cells("A1:G1")

headers2 = ["Rank", "Process", "Total Hours", "Avg Hours", "Avg Error %", "ROI Score", "Priority"]
for col, header in enumerate(headers2, 1):
    cell = ws2.cell(row=2, column=col, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="16213E")

for row_idx, (_, row) in enumerate(analysis.iterrows(), 3):
    ws2.cell(row=row_idx, column=1, value=int(row["rank"]))
    ws2.cell(row=row_idx, column=2, value=row["process"])
    ws2.cell(row=row_idx, column=3, value=round(row["total_manual_hours"], 2))
    ws2.cell(row=row_idx, column=4, value=round(row["avg_manual_hours"], 2))
    ws2.cell(row=row_idx, column=5, value=round(row["avg_error_rate"], 2))
    ws2.cell(row=row_idx, column=6, value=row["roi_score"])
    ws2.cell(row=row_idx, column=7, value=row["automation_priority"])

for col in ["A","B","C","D","E","F","G"]:
    ws2.column_dimensions[col].width = 22

# --- SHEET 3: RECOMMENDATIONS ---
ws3 = wb.create_sheet("Recommendations")

ws3["A1"] = "AI & Automation Recommendations by Process"
ws3["A1"].font = Font(bold=True, size=13, color="FFFFFF")
ws3["A1"].fill = PatternFill("solid", fgColor="0F3460")
ws3.merge_cells("A1:C1")

rec_headers = ["Process", "Recommended AI Solution", "Expected Benefit"]
for col, header in enumerate(rec_headers, 1):
    cell = ws3.cell(row=2, column=col, value=header)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="16213E")

recommendations = [
    ("Order Processing",        "Robotic Process Automation (RPA)",     "Reduce manual time by 80%"),
    ("Inventory Management",    "AI-based demand forecasting",           "Cut stockouts by 60%"),
    ("Customer Support",        "AI Chatbot (NLP)",                      "Handle 70% queries automatically"),
    ("Returns & Refunds",       "Automated refund approval system",      "Reduce processing time by 75%"),
    ("Fraud Detection",         "ML anomaly detection model",            "Reduce fraud errors by 90%"),
    ("Product Recommendations", "Collaborative filtering algorithm",     "Increase sales by 35%"),
    ("Delivery Tracking",       "IoT + automated status updates",        "Eliminate manual tracking"),
    ("Seller Onboarding",       "Automated KYC verification system",     "Reduce onboarding time by 50%"),
    ("Payment Reconciliation",  "Automated matching algorithm",          "99% accuracy, zero manual effort"),
    ("Review Moderation",       "NLP-based sentiment classifier",        "Moderate 10x faster"),
]

for row_idx, (process, solution, benefit) in enumerate(recommendations, 3):
    ws3.cell(row=row_idx, column=1, value=process)
    ws3.cell(row=row_idx, column=2, value=solution)
    ws3.cell(row=row_idx, column=3, value=benefit)

ws3.column_dimensions["A"].width = 28
ws3.column_dimensions["B"].width = 35
ws3.column_dimensions["C"].width = 35

# Save the workbook
os.makedirs("output", exist_ok=True)
wb.save("output/transformation_report.xlsx")
print("Excel report saved to output/transformation_report.xlsx")
print("")

# ============================================
# STEP 4 — CREATE DASHBOARD CHARTS
# ============================================

print("Creating dashboard charts...")

os.makedirs("dashboard", exist_ok=True)

fig, axes = plt.subplots(1, 3, figsize=(18, 6))
fig.suptitle("E-Commerce Business Transformation Dashboard", fontsize=16, fontweight="bold", color="#0f3460")

# Chart 1 — Manual Hours by Process
colors1 = ["#0f3460" if x == "HIGH" else "#16213e" if x == "MEDIUM" else "#a0aec0"
           for x in analysis["automation_priority"]]

axes[0].barh(analysis["process"], analysis["avg_manual_hours"], color=colors1)
axes[0].set_title("Avg Manual Hours by Process", fontweight="bold", color="#0f3460")
axes[0].set_xlabel("Hours")
axes[0].invert_yaxis()

# Chart 2 — Error Rate by Process
colors2 = ["#e53e3e" if x >= 20 else "#ed8936" if x >= 10 else "#48bb78"
           for x in analysis["avg_error_rate"]]

axes[1].barh(analysis["process"], analysis["avg_error_rate"], color=colors2)
axes[1].set_title("Avg Error Rate % by Process", fontweight="bold", color="#0f3460")
axes[1].set_xlabel("Error Rate %")
axes[1].invert_yaxis()

# Chart 3 — ROI Score (Automation Priority)
colors3 = ["#0f3460" if x == "HIGH" else "#16213e" if x == "MEDIUM" else "#a0aec0"
           for x in analysis["automation_priority"]]

bars = axes[2].bar(analysis["process"], analysis["roi_score"], color=colors3)
axes[2].set_title("ROI Score — Automation Priority", fontweight="bold", color="#0f3460")
axes[2].set_ylabel("ROI Score")
axes[2].set_xticklabels(analysis["process"], rotation=45, ha="right", fontsize=8)

# Legend
from matplotlib.patches import Patch
legend = [
    Patch(color="#0f3460", label="HIGH Priority"),
    Patch(color="#16213e", label="MEDIUM Priority"),
    Patch(color="#a0aec0", label="LOW Priority"),
]
fig.legend(handles=legend, loc="lower center", ncol=3, fontsize=10)

plt.tight_layout(rect=[0, 0.05, 1, 1])
plt.savefig("dashboard/transformation_dashboard.png", dpi=150, bbox_inches="tight")
plt.show()

print("Dashboard saved to dashboard/transformation_dashboard.png")
print("")
print("============================================")
print("PROJECT COMPLETE!")
print("============================================")
print("Files generated:")
print("  output/transformation_report.xlsx")
print("  dashboard/transformation_dashboard.png")
print("============================================")